from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT=Path('output/docx/Arquitetura_Completa_Piloto_WhatsApp_Life.docx'); OUT.parent.mkdir(parents=True,exist_ok=True)
CHAR='252525'; INK='353535'; TERR='A34F38'; SAND='F3EEE8'; PALE='FCFAF7'; LINE='D9D0C7'; WHITE='FFFFFF'
def shade(c, fill):
 p=c._tc.get_or_add_tcPr(); e=OxmlElement('w:shd'); e.set(qn('w:fill'),fill); p.append(e)
def borders(c):
 p=c._tc.get_or_add_tcPr(); x=OxmlElement('w:tcBorders')
 for n in ('top','left','bottom','right'):
  e=OxmlElement('w:'+n); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:color'),LINE); x.append(e)
 p.append(x)
def setrun(r,size=10.2,bold=False,color=INK,font='Liberation Sans'):
 r.font.name=font; r._element.rPr.rFonts.set(qn('w:eastAsia'),font); r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=RGBColor.from_string(color)
def celltext(c,v,size=8.6,bold=False,color=INK):
 p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.08
 r=p.add_run(v); setrun(r,size,bold,color)
def page_num(p):
 p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; r=p.add_run('Life Engenharia  |  Arquitetura do piloto  |  '); setrun(r,8,False,'777777')
 b=OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'),'begin'); i=OxmlElement('w:instrText'); i.set(qn('xml:space'),'preserve'); i.text='PAGE'; e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'end'); r._r.extend([b,i,e])

d=Document(); s=d.sections[0]; s.top_margin=Inches(.7);s.bottom_margin=Inches(.65);s.left_margin=Inches(.72);s.right_margin=Inches(.72);page_num(s.footer.paragraphs[0])
styles=d.styles; n=styles['Normal'];n.font.name='Liberation Sans';n._element.rPr.rFonts.set(qn('w:eastAsia'),'Liberation Sans');n.font.size=Pt(10.2);n.font.color.rgb=RGBColor.from_string(INK);n.paragraph_format.space_after=Pt(6);n.paragraph_format.line_spacing=1.14
for st,sz,col,font in [('Title',27,CHAR,'Liberation Serif'),('Subtitle',12,'676767','Liberation Sans'),('Heading 1',18,CHAR,'Liberation Serif'),('Heading 2',12,TERR,'Liberation Sans')]:
 x=styles[st];x.font.name=font;x._element.rPr.rFonts.set(qn('w:eastAsia'),font);x.font.size=Pt(sz);x.font.color.rgb=RGBColor.from_string(col);x.font.bold=st!='Subtitle';x.paragraph_format.space_before=Pt(14 if st.startswith('Heading') else 0);x.paragraph_format.space_after=Pt(7)
def p(t='',style=None):
 q=d.add_paragraph(style=style);r=q.add_run(t);setrun(r,10.2,False,INK);return q
def h(t,l=1):return d.add_paragraph(t,style='Heading '+str(l))
def bullets(xs):
 for x in xs:
  q=d.add_paragraph(style='List Bullet');q.paragraph_format.space_after=Pt(3);r=q.add_run(x);setrun(r,10)
def table(head,rows,widths=None):
 t=d.add_table(rows=1,cols=len(head));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,v in enumerate(head):
  c=t.rows[0].cells[i];shade(c,CHAR);borders(c);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;celltext(c,v,8.5,True,WHITE); 
  if widths:c.width=Inches(widths[i])
 for ri,row in enumerate(rows):
  cells=t.add_row().cells
  for i,v in enumerate(row):
   if ri%2:shade(cells[i],PALE)
   borders(cells[i]);cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;celltext(cells[i],str(v));
   if widths:cells[i].width=Inches(widths[i])
 d.add_paragraph().paragraph_format.space_after=Pt(2)
 return t
def br():d.add_page_break()

# Cover
d.add_paragraph(); q=d.add_paragraph();q.paragraph_format.space_after=Pt(16);r=q.add_run('LIFE ENGENHARIA  •  DOCUMENTO TÉCNICO CONSOLIDADO');setrun(r,9,True,TERR)
d.add_paragraph('Arquitetura completa do piloto de atendimento por WhatsApp',style='Title')
d.add_paragraph('Referência técnica para homologação, desenvolvimento e operação controlada.',style='Subtitle')
d.add_paragraph().paragraph_format.space_after=Pt(25)
table(['OBJETIVO','DECISÃO DE ARQUITETURA'],[['Automatizar a triagem e o roteamento, mantendo decisões técnicas, prioridades críticas e atendimento humano sob responsabilidade da Life.','Meta WhatsApp Cloud API + Node.js/Fastify no Google Cloud Run + Supabase Pro/PostgreSQL + OpenAI API + Ploomes comercial.']],[3.18,3.18])
d.add_paragraph().paragraph_format.space_after=Pt(60)
p('Preparado por João Victor Vasconcelos  |  Engenharia de Software')
p('Versão consolidada  |  Setembro de 2026')
br()

h('1  Visão geral e escopo')
p('O piloto é um orquestrador de triagem. Ele não substitui a decisão da equipe técnica, não oferece diagnóstico, não define risco final e não promete prazo. Seu papel é apresentar opções, coletar os dados necessários, gerar protocolo e resumo, sugerir prioridade por regras e encaminhar ao responsável humano.')
table(['Incluído no piloto','Fora do piloto'],[
['Canal oficial WhatsApp; rotas comercial, assistência e pessoa; triagem; mídia; protocolo; escalonamento P1/P2; operação interna privada; Ploomes comercial.','Dashboard gerencial, BI, integração com sistema de assistência técnica, diagnóstico automatizado, automações de agenda ou financeiro.']
],[3.18,3.18])
h('Princípios de projeto',2)
bullets(['API oficial da Meta; não utilizar automação de WhatsApp Web.','Fluxo determinístico para menu, campos obrigatórios, estado e escalonamento.','IA como apoio para compreender, extrair e resumir; nunca como decisão de risco.','Dados mínimos, acessos por perfil e todas as ações relevantes auditáveis.','Infraestrutura gerenciada e escalável sem Kubernetes ou microserviços no piloto.'])
h('Público e evolução',2)
p('A liberação inicial é para 5 a 10 clientes selecionados, acompanhados diariamente. A estimativa de 100 clientes é usada apenas para orientar capacidade, custos e evolução de infraestrutura.')
br()

h('2  Arquitetura lógica')
p('O fluxo é dividido entre o canal do cliente, o orquestrador, os serviços de dados e as conversas privadas da equipe. O cliente nunca visualiza o número pessoal do colaborador.')
table(['Camada','Componente','Responsabilidade'],[
['Canal externo','Meta WhatsApp Cloud API e Webhooks','Receber mensagens, mídia e eventos; enviar resposta pelo número oficial da Life.'],
['Entrada','Webhook Fastify','Validar assinatura Meta, identificar duplicidade, registrar evento e responder rapidamente.'],
['Orquestração','Node.js, TypeScript e Fastify','Executar máquina de estados, regras, roteamento, protocolo, limites e contingência.'],
['Dados','Supabase Pro com PostgreSQL','Persistir contatos, conversas, mensagens, protocolos, equipe, escalas, auditoria e integrações.'],
['Arquivos','Supabase Storage e cópia no Google Cloud Storage','Guardar anexos privados e cópias de backup cifradas.'],
['IA','OpenAI API','Interpretar texto, extrair campos, resumir e processar mídia dentro das regras aprovadas.'],
['CRM','Ploomes API','Criar ou atualizar contatos e oportunidades somente na trilha comercial.'],
['Operação interna','Bot privado no WhatsApp','Notificar, confirmar, assumir, transferir e registrar ações por perfil autorizado.']
],[.85,1.75,3.75])
h('Diagrama de fluxo',2)
p('Cliente → WhatsApp oficial Life → Meta Webhook → Fastify no Cloud Run → regras e máquina de estados → Supabase / OpenAI / Ploomes → bot interno e responsável autorizado → resposta pelo número oficial Life.')
br()

h('3  Jornada do cliente e máquina de estados')
table(['Estado','Regra de transição'],[
['NOVO → CONSENTIMENTO → ESCOLHA ROTA','Apresenta o assistente virtual, registra consentimento quando aplicável e mostra o menu inicial.'],
['CONTRATAÇÃO COLETA → RESUMO COMERCIAL → FILA COMERCIAL','Coleta serviço, contato e necessidade; pede confirmação; cria oportunidade comercial e notifica Comercial.'],
['CLIENTE LIFE COLETA → VALIDAR CONTRATO → TÉCNICA COLETA','Confirma contrato ou histórico; coleta equipamento, sintoma, impacto, contingência e evidências.'],
['PRIORIDADE → FILA TÉCNICA → HUMANO','Regras sugerem P1 a P4; humano confirma P1/P2; responsável assume e responde.'],
['HANDOFF HUMANO → ENCERRADA','Quando houver pedido de pessoa, urgência, baixa confiança, falha de integração, reclamação grave ou direito LGPD.']
],[2.15,4.2])
h('Menu inicial',2)
table(['Opção','Resultado'],[
['1  Contratar serviço ou solicitar orçamento','Submenu comercial e criação/atualização de oportunidade no Ploomes.'],
['2  Já sou cliente Life e preciso de assistência técnica','Pré-triagem e encaminhamento para fila humana técnica.'],
['3  Falar com uma pessoa','Protocolo e transferência imediata para a fila compatível.']
],[2.25,4.1])
h('Submenu comercial',2)
p('Limpeza de dutos; PMOC; assistência técnica avulsa; projetos de engenharia; Engenharia Clínica; locação de chillers; venda de chiller novo; Banco de Sangue; calibração de equipamentos e instrumentos de medição; outros serviços.')
p('Solicitação de assistência sem contrato ou sem histórico Life não é recusada: segue como oportunidade comercial de assistência avulsa, preservando os dados já coletados.')
br()

h('4  Operação interna, papéis e escalonamento')
p('A interface operacional do piloto é um bot privado no WhatsApp. Somente números corporativos ou explicitamente autorizados podem usar os comandos internos. Não há dashboard gerencial nesta fase.')
table(['Perfil','Permissões principais'],[
['Comercial','Assumir oportunidades, consultar triagem comercial, responder cliente, atualizar oportunidade no Ploomes e transferir.'],
['Assistência técnica','Assumir chamado, acessar histórico e anexos autorizados, responder pelo número oficial, registrar ação e transferir.'],
['Plantão','Confirmar e assumir P1/P2 em cobertura, acionar backup e registrar ações.'],
['Backup','Receber somente casos escalados ou encaminhados pela Coordenação; assumir ou declarar indisponibilidade.'],
['Coordenação','Acompanhar fila, designar responsável, ajustar prioridade humana, transferir, revisar histórico e encerrar.'],
['Administração','Cadastrar números, cargos, escalas, backups e permissões; consultar auditoria e acionar LGPD/segurança.']
],[1.45,4.9])
h('Escalonamento P1 e P2',2)
table(['Prioridade','Fluxo proposto'],[
['P1','Notifica responsável principal. Sem confirmação em até 5 minutos após SLA aprovado, notifica backup. Sem confirmação, aciona Coordenação/Plantão.'],
['P2','Notifica responsável principal. Sem confirmação em até 15 minutos após SLA aprovado, notifica backup e registra evento.'],
['P3, P4 e Comercial','Distribuição conforme equipe, horário e regra operacional aprovada.']
],[1.1,5.25])
p('A Life deve definir nominalmente principal, backup, telefone, cargo, horário, escala e contingência antes da ativação. Cada notificação, confirmação e transferência gera evento de auditoria.')
br()

h('5  Processamento de mensagens e mídia')
table(['Entrada','Tratamento permitido no piloto'],[
['Texto','Classificação de intenção, coleta de campos, resposta aprovada e resumo.'],
['Imagem ou foto','Leitura visual para contexto da triagem; arquivo original preservado. Não há diagnóstico automático.'],
['Áudio','Transcrição para contexto e resumo; áudio original preservado.'],
['PDF, documento ou planilha','Extração de texto e leitura visual quando necessária; conteúdo sensível encaminhado ao humano.'],
['Vídeo','Extração de áudio e quadros selecionados para triagem; vídeo original preservado.'],
['Arquivo ilegível, protegido ou fora do limite','Marcar para análise humana; nunca inventar conteúdo.']
],[1.55,4.8])
h('Guardrails de IA',2)
bullets(['Enviar à IA somente contexto mínimo e versão vigente das regras.','Não instruir operação elétrica, mecânica, clínica ou sanitária; não diagnosticar equipamento.','Não prometer preço, estoque, agenda, prazo ou disponibilidade sem fonte autorizada.','Exigir confirmação antes de criar oportunidade ou enviar dados relevantes.','Acionar humano quando houver risco, ambiguidade, baixa confiança ou solicitação expressa.'])
h('Ferramenta de apoio à equipe',2)
p('Claude Pro ou Claude Team são opcionais e independentes do bot. Servem apenas para desenvolvimento, revisão de código, documentação, análise de erros, testes e organização técnica. Não recebem dados reais de clientes, conversas ou anexos.')
br()

h('6  Dados, banco e integrações')
h('Banco definido: Supabase Pro com PostgreSQL',2)
p('PostgreSQL foi escolhido pela consistência das relações entre cliente, unidade, protocolo, prioridade, responsável, backup, mensagens e auditoria. Um banco NoSQL puro não traz vantagem no piloto; poderá ser adicionado apenas como cache ou fila no crescimento.')
table(['Entidade','Campos essenciais'],[
['contacts','telefone, nome, empresa, e-mail opcional, consentimento e timestamps.'],
['conversations','telefone, rota, estado, protocolo, prioridade, responsável e timestamps.'],
['messages','ID Meta, conversa, direção, conteúdo, mídia, status de entrega e idempotência.'],
['commercial_requests','serviço, local, prazo, necessidade, Ploomes ID e status.'],
['technical_tickets','contrato, equipamento, modelo, sintoma, impacto, contingência, risco e prioridade sugerida/confirmada.'],
['handoffs e audit_events','destino, confirmação, SLA, usuário, ação, responsável anterior/novo e resultado.'],
['team_members e schedules','telefone autorizado, perfil, equipe, horário, plantão, principal, backup e situação.']
],[1.85,4.5])
h('Ploomes confirmado',2)
p('A API do Ploomes está incluída no plano atual da Life, sem custo adicional de uso. A Life terá usuário exclusivo de integração e User-Key guardada no Secret Manager. O perfil será limitado a Contatos e Oportunidades. Não existe ambiente de teste dedicado no Ploomes; por isso, testes deverão usar registros controlados, identificados e reversíveis.')
table(['Recurso Ploomes','Uso'],[['Contatos','Localizar por telefone ou e-mail; criar ou atualizar dados confirmados pelo cliente.'],['Oportunidades','Criar ou atualizar serviço, unidade, necessidade, prazo, protocolo Life e responsável comercial.'],['Catálogos','Consultar funis, estágios, usuários, equipes, campos e tags.'],['Falha','Manter protocolo local, registrar tentativa e alertar o Comercial; reprocessamento manual.']],[1.45,4.9])
br()

h('7  Deploy, ambientes e observabilidade')
table(['Ambiente','Configuração'],[
['Homologação','Projeto/serviço isolado, credenciais próprias, dados fictícios ou anonimizados e número Meta de teste quando disponível.'],
['Produção','Cloud Run na região southamerica-east1, São Paulo; mínimo de zero instâncias; acesso restrito e logs monitorados.'],
['CI/CD','Repositório GitHub da Life, branches protegidas, revisão de código, Docker e deploy com conta de serviço de menor privilégio.'],
['Segredos','Meta token, OpenAI key, Ploomes User-Key e conexão de serviços somente no Google Secret Manager.'],
['Observabilidade','Logs estruturados por protocolo, alertas de falha, monitoramento de webhook, consumo de APIs e P1/P2 sem confirmação.']
],[1.25,5.1])
h('Google Cloud sob controle institucional',2)
p('O projeto Google Cloud é da Life. A configuração final inclui faturamento institucional, pelo menos dois administradores institucionais, acesso individual por função, MFA, orçamento com alertas e revisão periódica de permissões. João recebe acesso técnico de menor privilégio, revogável e nunca a propriedade exclusiva da conta.')
h('Evolução sem reescrita',2)
table(['Gatilho','Evolução'],[['Mais mensagens simultâneas','Adicionar Redis gerenciado e BullMQ para desacoplar webhook e tarefas.'],['Conteúdo amplo','Adicionar busca vetorial/RAG, mantendo PostgreSQL como fonte de verdade.'],['Maior criticidade','Ampliar observabilidade, cópias de backup, ambientes e revisão de segurança.'],['Dashboard ou BI','Tratar como fase posterior, consumindo dados já auditados do PostgreSQL.']],[2.0,4.35])
br()

h('8  Segurança, backup e LGPD')
h('Controles obrigatórios',2)
bullets(['Validação de assinatura de webhooks Meta, idempotência por ID de mensagem e proteção contra loops.','MFA em administradores e serviços críticos; credenciais individuais e menor privilégio.','Números não cadastrados não recebem protocolo, anexos ou comandos operacionais.','Criptografia em trânsito e repouso; anexos privados e links temporários quando necessários.','Auditoria imutável de mudança de status, prioridade, responsável, escalonamento e envio ao cliente.','Separação estrita entre homologação e produção.'])
h('Política inicial de backup',2)
table(['Componente','Regra de referência'],[['Banco','Backup diário Supabase Pro, retenção nativa de 7 dias e exportação cifrada semanal para bucket privado da Life com retenção de 30 dias.'],['Anexos','Cópia cifrada diária para bucket privado; retenção de cópia inicialmente em 30 dias.'],['Restauração','Teste antes da produção e mensal durante o piloto. Meta inicial: RPO até 24h e RTO até 4h.'],['Falhas','Alertar backup ausente, cópia falha ou teste de restauração com erro.']],[1.25,5.1])
h('LGPD e dados sensíveis',2)
p('A Life é a controladora dos dados. Administração e Jurídico devem validar base legal, aviso de privacidade, retenção, descarte, atendimento aos titulares, resposta a incidentes e contratos com operadores antes da produção. Dados de pacientes, doadores, prontuários e outros dados sensíveis não devem ser solicitados pelo fluxo padrão; se recebidos, o acesso é limitado e o caso é direcionado ao humano.')
table(['Dado','Retenção inicial a validar'],[['Conversas e protocolos','180 dias, salvo obrigação legal ou contratual.'],['Anexos','90 dias, salvo necessidade técnica, operacional ou contratual.'],['Backups','30 dias para cópias externas, conforme política aprovada.']],[2.1,4.25])
br()

h('9  Custos, critérios de aceite e pendências')
p('Valores são estimativas de referência e devem ser validados antes de qualquer contratação. Não incluem impostos, IOF, desenvolvimento, suporte recorrente ou gastos fora do escopo.')
table(['Componente','Piloto 5 a 10 clientes','Projeção 100 clientes'],[['Supabase Pro produção','R$ 137,50','R$ 137,50 ou mais'],['Homologação Supabase','R$ 0 a R$ 55','R$ 55'],['Cloud Run, Registry, segredos, logs e backup','R$ 2 a R$ 30','R$ 10 a R$ 60'],['OpenAI API multimodal','R$ 5 a R$ 15','R$ 30 a R$ 90'],['Meta WhatsApp Cloud API','R$ 5 a R$ 20','R$ 70 ou mais'],['GitHub Team, se necessário','R$ 66 para 3 usuários','R$ 110 para 5 usuários'],['Ploomes API','Sem custo adicional confirmado','Sem custo adicional confirmado'],['Total técnico estimado','R$ 215 a R$ 324 por mês','R$ 410 a R$ 560 por mês']],[2.8,1.75,1.75])
h('Critérios para liberar o piloto',2)
bullets(['Responsáveis, backups, telefones, escalas e horários cadastrados e testados.','P1/P2 com confirmação e escalonamento validados em cenários de teste.','Ploomes com credencial restrita, idempotência e contingência testadas.','Backups, restauração, auditoria, alertas e acesso por perfil revisados.','Pelo menos 20 cenários aprovados, incluindo mídia, duplicidade, falha de integração, assistência sem contrato e pedido de humano.','Aprovação da Administração e Jurídico sobre LGPD antes de produção.'])
h('Pendências de decisão da Life',2)
table(['Tema','Decisão necessária'],[['Operação','Nome, telefone, cargo, principal, backup, horário e contingência por equipe/turno.'],['SLA','Confirmação formal de P1 em 5 minutos e P2 em 15 minutos, ou ajuste dos prazos.'],['LGPD','Retenção, dados sensíveis, aviso de privacidade, canal de titulares e resposta a incidentes.'],['Contas','Confirmar faturamento, administradores institucionais e alertas nas contas Meta, Google Cloud, Supabase, GitHub e OpenAI.']],[1.4,4.95])

d.core_properties.author='João Victor Vasconcelos';d.core_properties.title='Arquitetura Completa Piloto WhatsApp Life';d.save(OUT);print(OUT)
