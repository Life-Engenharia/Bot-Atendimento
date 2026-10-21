from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'approvals' / 'MINUTA_POLITICA_PRIVACIDADE_CENTRAL_ATENDIMENTO_LIFE.docx'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); b.append(e)
    tblPr.append(b)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.75); sec.bottom_margin = Inches(.7); sec.left_margin = Inches(.85); sec.right_margin = Inches(.85)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10.5); styles['Normal'].font.color.rgb = RGBColor(35,45,55)
styles['Normal'].paragraph_format.space_after = Pt(6)
for name, size, color in [('Title',22,'0B3553'),('Heading 1',15,'0B3553'),('Heading 2',12,'176B9A')]:
    st=styles[name]; st.font.name='Aptos Display'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color); st.paragraph_format.space_before=Pt(14); st.paragraph_format.space_after=Pt(6)

title = doc.add_paragraph(style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.LEFT; title.add_run('Política de privacidade e governança de dados')
sub = doc.add_paragraph(); sub.add_run('Central de Atendimento Life Engenharia').bold = True; sub.runs[0].font.color.rgb = RGBColor(23,107,154)
meta = doc.add_paragraph(); meta.add_run('MINUTA PARA REVISÃO JURÍDICA\n').bold=True; meta.add_run('Versão 0.1 · 21/10/2026 · Não aprovada')
doc.add_paragraph('Esta minuta descreve o funcionamento planejado da Central de Atendimento. O Jurídico deve validar a identidade da controladora, bases legais, prazos, contratos com fornecedores, transferências internacionais e textos finais antes da utilização com dados reais.')

sections = [
('1. Identificação e escopo', ['Esta política se aplica à Central de Atendimento da Life Engenharia (“Life”), utilizada para receber solicitações pelo WhatsApp oficial, organizar triagens, encaminhar atendimentos às equipes e acompanhar a resolução.','O documento cobre mensagens, dados cadastrais, informações técnicas, anexos, registros de atendimento, notificações internas e dados dos operadores.']),
('2. Papéis no tratamento', ['• Controladora: Life Engenharia, a confirmar pelo Jurídico.','• Operadores e fornecedores: provedores contratados para infraestrutura, mensageria, armazenamento, inteligência artificial, CRM, e-mail e hospedagem.','• Usuários internos: colaboradores autorizados por área, cargo e necessidade operacional.','• Titulares: clientes, representantes, funcionários de clientes, prestadores e demais pessoas que enviem mensagens ou apareçam nos dados do atendimento.']),
('3. Dados tratados', ['3.1 Dados fornecidos pelo solicitante','• nome e telefone; • e-mail, quando informado; • empresa, instituição, cidade e unidade; • contrato ou histórico informado; • equipamento, patrimônio, número de série e sintomas; • prazo, impacto, risco, contingência e descrição da necessidade; • documentos, imagens, vídeos e áudios enviados voluntariamente.','3.2 Dados gerados pela operação','• protocolo, rota, prioridade, fila e status; • mensagens recebidas e enviadas; • responsável, backup, horários, transferências e confirmações; • registros de entrega, leitura e falha; • auditoria, logs técnicos e tentativas de integração.','3.3 Dados potencialmente sensíveis','A operação pode receber informações que revelem risco à saúde, funcionamento de hospitais, sangue, medicamentos ou pessoas. O bot deve solicitar somente o mínimo necessário, não deve pedir diagnóstico médico e não deve orientar ações inseguras.']),
('4. Finalidades', ['Os dados são tratados para receber e registrar solicitações; confirmar autorização do atendimento; realizar triagem; encaminhar casos; permitir que operadores assumam, respondam e concluam; criar protocolos e oportunidades; enviar notificações; disponibilizar anexos autorizados; melhorar segurança e auditoria; e cumprir obrigações legais, regulatórias e contratuais.','Não usar os dados para publicidade, venda de listas ou treinamento genérico de modelos sem base legal, transparência e aprovação específica.']),
('5. Bases legais a validar', ['O Jurídico deve confirmar, finalidade por finalidade, a base legal aplicável, que pode incluir execução de contrato ou procedimentos preliminares, legítimo interesse, cumprimento de obrigação legal ou regulatória, exercício regular de direitos e consentimento quando necessário.','O consentimento do fluxo conversacional não substitui a análise jurídica da base legal.']),
('6. Compartilhamento e fornecedores', ['Meta / WhatsApp Cloud API: recebe e envia mensagens pelo número oficial da Life.','Supabase: armazena dados relacionais, mensagens, auditoria e anexos privados.','OpenAI: pode processar triagem, transcrição e leitura/resumo quando habilitado.','Ploomes: pode receber dados mínimos para lead ou oportunidade comercial.','Google Cloud e e-mail corporativo: podem hospedar aplicação, logs e notificações.','O Jurídico deve avaliar termos, localização, transferência internacional, retenção e contratos de cada fornecedor.']),
('7. Segurança', ['A Life deve manter HTTPS; tokens em Secret Manager; rotação de credenciais; autenticação por função; menor privilégio; assinatura dos webhooks; deduplicação; URLs temporárias; backups; logs sem conteúdo sensível desnecessário; e revisão de acessos após mudança ou desligamento.']),
('8. Retenção proposta', ['Os prazos abaixo são propostas para aprovação e devem considerar obrigação legal, necessidade contratual, defesa de direitos e backup.']),
('9. Direitos dos titulares', ['O canal definido pelo Jurídico deverá permitir solicitações de confirmação de tratamento, acesso, correção, anonimização, bloqueio, eliminação, portabilidade quando aplicável, informação sobre compartilhamentos e revogação de consentimento.','Solicitações devem ser autenticadas, registradas, encaminhadas ao responsável e respondidas dentro dos prazos legais aplicáveis.']),
('10. Incidentes de segurança', ['Qualquer suspeita de acesso indevido, vazamento, perda, indisponibilidade ou envio incorreto deve ser registrada imediatamente. O procedimento deve conter contenção, preservação de evidências, avaliação de risco, identificação dos titulares afetados, comunicação interna e decisão sobre comunicação à ANPD e aos titulares.']),
('11. Governança e aprovação', ['• Dados do controlador preenchidos\n• Encarregado/DPO e canal definidos\n• Bases legais validadas\n• Contratos dos fornecedores revisados\n• Transferências internacionais avaliadas\n• Prazos de retenção aprovados\n• Procedimento de direitos aprovado\n• Plano de incidentes aprovado\n• Backup e restauração testados\n• Política publicada e versão registrada'])]

for heading, paras in sections:
    doc.add_heading(heading, level=1)
    for text in paras:
        if text.startswith('3.'):
            doc.add_heading(text, level=2)
        elif text.startswith('•'):
            for item in text.split('• '):
                if item.strip(): doc.add_paragraph(item.strip(), style='List Bullet')
        else: doc.add_paragraph(text)

doc.add_heading('Tabela de retenção proposta', level=1)
table = doc.add_table(rows=1, cols=3); table.style='Table Grid'; borders(table)
for i, h in enumerate(['Registro','Prazo proposto','Ação ao vencer']):
    cell=table.rows[0].cells[i]; cell.text=h; shade(cell,'0B3553')
    for run in cell.paragraphs[0].runs: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255)
rows=[('Conversa e protocolo','180 dias após encerramento','Anonimizar ou excluir'),('Anexos','90 dias após encerramento','Excluir do Storage privado'),('Logs de auditoria','365 dias','Revisar e eliminar ou anonimizar'),('Backups','Política corporativa','Expirar com evidência'),('Credenciais e tokens','Até substituição/revogação','Revogar imediatamente se comprometidos')]
for row in rows:
    cells=table.add_row().cells
    for i,val in enumerate(row): cells[i].text=val

doc.add_paragraph('Campos para preenchimento pelo Jurídico', style='Heading 1')
for label in ['Razão social e CNPJ: __________________________________________','Endereço: _________________________________________________','Encarregado/DPO e canal: ____________________________________','Canal para solicitações de titulares: __________________________','Responsável por incidentes: _________________________________','Aprovador, cargo e data: ____________________________________']:
    doc.add_paragraph(label)

footer = sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.add_run('Life Engenharia · Minuta para revisão jurídica · Versão 0.1').font.size=Pt(8)
doc.core_properties.title = 'Política de privacidade e governança de dados — Central Life'
doc.core_properties.subject = 'Minuta para revisão jurídica'
doc.core_properties.author = 'Life Engenharia'
doc.save(OUT)
print(OUT)
